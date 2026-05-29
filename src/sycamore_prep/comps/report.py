"""Writers for the comps outputs: a multi-tab xlsx and a markdown tear-sheet.

xlsx reuses the screener's openpyxl conventions (bold frozen header, auto-width,
NaN/Inf -> blank). The markdown is hand-rolled (no `tabulate` dependency). Every
figure is source-tagged; the three lenses stay on separate tabs/sections so they
are never collapsed (CLAUDE.md #2).
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd

from .comps import CompsResult
from .history import HistoryResult
from .normalized import NormalizedEarnings


# --------------------------------------------------------------------------- #
# Shared cell / number formatting
# --------------------------------------------------------------------------- #

def _safe(v):
    """Coerce a value to something openpyxl can store; NaN/Inf -> None (blank)."""
    if isinstance(v, np.generic):
        v = v.item()
    if v is None:
        return None
    if isinstance(v, float) and (pd.isna(v) or np.isinf(v)):
        return None
    if isinstance(v, (int, float, str, bool)):
        return v
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    return str(v)


def _f(x, dp: int = 2) -> str:
    try:
        if x is None or (isinstance(x, float) and (pd.isna(x) or np.isinf(x))):
            return "n/a"
    except (TypeError, ValueError):
        return "n/a"
    if isinstance(x, (int, np.integer)):
        return f"{int(x):,}"
    return f"{float(x):,.{dp}f}"


def _pct(x, dp: int = 1) -> str:
    try:
        if x is None or (isinstance(x, float) and (pd.isna(x) or np.isinf(x))):
            return "n/a"
    except (TypeError, ValueError):
        return "n/a"
    return f"{float(x) * 100:.{dp}f}%"


def _money_m(x) -> str:
    try:
        if x is None or (isinstance(x, float) and (pd.isna(x) or np.isinf(x))):
            return "n/a"
    except (TypeError, ValueError):
        return "n/a"
    return f"{float(x) / 1e6:,.0f}"


# --------------------------------------------------------------------------- #
# Tab DataFrames
# --------------------------------------------------------------------------- #

def _history_detail_df(hist: HistoryResult | None) -> pd.DataFrame:
    if hist is None:
        return pd.DataFrame()
    cols: dict[str, pd.Series] = {"fy_end_price": hist.fy_end_prices,
                                  "market_cap": hist.market_caps}
    for name, band in hist.bands.items():
        cols[name] = band.series
    df = pd.DataFrame(cols)
    df.index.name = "period"
    return df.sort_index()


def _band_summary_df(hist: HistoryResult | None) -> pd.DataFrame:
    if hist is None:
        return pd.DataFrame()
    rows = [{
        "multiple": name,
        "current": b.current,
        "pctile_cheap_vs_own_history": b.percentile_cheap,
        "n": b.n,
        "min": b.p_min, "p25": b.p25, "median": b.median, "p75": b.p75, "max": b.p_max,
    } for name, b in hist.bands.items()]
    return pd.DataFrame(rows).set_index("multiple")


def _normalized_df(norm: NormalizedEarnings | None) -> pd.DataFrame:
    if norm is None:
        return pd.DataFrame()
    items = {
        "window_years": norm.window,
        "n_used": norm.n_used,
        "mean_op_margin": norm.mean_op_margin,
        "trough_op_margin": norm.trough_op_margin,
        "current_revenue": norm.current_revenue,
        "interest_expense": norm.interest_expense,
        "tax_rate": norm.tax_rate,
        "normalized_op_income": norm.normalized_op_income,
        "normalized_net_income": norm.normalized_net_income,
        "normalized_eps": norm.normalized_eps,
        "trough_net_income": norm.trough_net_income,
        "trough_eps": norm.trough_eps,
        "trailing_pe": norm.trailing_pe,
        "normalized_pe": norm.normalized_pe,
        "trough_pe": norm.trough_pe,
        "basis": norm.basis,
    }
    df = pd.DataFrame({"value": items})
    df.index.name = "metric"
    return df


def _dcf_df(cases: list) -> pd.DataFrame:
    if not cases:
        return pd.DataFrame()
    rows = [{
        "case": c.label,
        "wacc": c.wacc,
        "terminal_growth": c.terminal_growth,
        "fcf0": c.fcf0,
        "implied_growth": c.implied_growth,
        "implied_converged": c.implied_converged,
        "assumed_growth": c.assumed_growth,
        "fair_value_per_share": c.fair_value_per_share,
        "current_price": c.current_price,
        "margin_of_safety": c.margin_of_safety,
    } for c in cases]
    return pd.DataFrame(rows).set_index("case")


def _sources_df(result: CompsResult) -> pd.DataFrame:
    s = result.subject
    items = {
        "subject": s.ticker,
        "subject_sources": s.sources,
        "share_basis_used": s.history.share_basis if s.history else "n/a",
        "market_cap": s.market_cap,
        "current_price": s.price,
        "wacc": result.wacc,
        "terminal_growth": result.terminal_growth,
        "forecast_years": result.forecast_years,
        "fundamentals": "SEC EDGAR XBRL (primary)",
        "prices_and_market_cap": "yfinance (non-primary)",
        "note_bands": "Own-history percentiles need >=3 annual observations; fewer -> NaN.",
        "note_prices": "Unadjusted FY-end close x as-reported shares = historical market cap.",
        "note_splits": "No general split adjustment; a split inside the lookback would distort a band.",
        "note_dcf": "Reverse DCF: FCFF, end-of-year discounting, Gordon terminal; bisection solver.",
        "note_bank": "Banks: P/E + P/TBV bands + normalized EPS; reverse DCF = N/A.",
    }
    df = pd.DataFrame({"value": items})
    df.index.name = "key"
    return df


# --------------------------------------------------------------------------- #
# xlsx
# --------------------------------------------------------------------------- #

def _df_to_sheet(ws, df: pd.DataFrame, index_label: str) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    if df is None or df.empty:
        ws.cell(row=1, column=1, value="(no data)")
        return
    cols = [index_label] + list(df.columns)
    for c, v in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=str(v))
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
        cell.alignment = Alignment(horizontal="left")
    for r, (idx, rowv) in enumerate(df.iterrows(), start=2):
        ws.cell(row=r, column=1, value=_safe(idx))
        for c, col in enumerate(df.columns, start=2):
            ws.cell(row=r, column=c, value=_safe(rowv[col]))
    ws.freeze_panes = "B2"
    for i, col in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(
            12, min(30, len(str(col)) + 4)
        )


def write_comps_xlsx(result: CompsResult, path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    sheets = [
        ("peer_comps", result.peer_table.reset_index().rename(columns={"index": "ticker"})
            .set_index("ticker"), "ticker"),
        ("history_bands", _history_detail_df(result.subject.history), "period"),
        ("band_summary", _band_summary_df(result.subject.history), "multiple"),
        ("normalized", _normalized_df(result.subject.normalized), "metric"),
        ("reverse_dcf", _dcf_df(result.subject.dcf_cases), "case"),
        ("sources", _sources_df(result), "key"),
    ]
    for i, (title, df, idx_label) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        _df_to_sheet(ws, df, idx_label)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# --------------------------------------------------------------------------- #
# markdown tear-sheet
# --------------------------------------------------------------------------- #

def _headline_read(result: CompsResult) -> str:
    s = result.subject
    parts: list[str] = []
    pe = s.history.bands.get("PE") if s.history else None
    if pe is not None and pe.percentile_cheap == pe.percentile_cheap:
        parts.append(
            f"P/E sits at the **{pe.percentile_cheap:.0f}th** cheapness percentile of its "
            f"own history (100 = cheapest vs the last {pe.n} FYs)."
        )
    g = s.base_implied_growth()
    if g == g:
        parts.append(f"Base reverse DCF implies **{g * 100:.1f}%** near-term FCFF growth priced in.")
    mos = s.base_margin_of_safety()
    if mos == mos:
        parts.append(f"Base-case margin of safety: **{mos * 100:.1f}%**.")
    return " ".join(parts) if parts else "_Insufficient data for a headline read._"


def write_comps_md(result: CompsResult, path: Path) -> None:
    s = result.subject
    L: list[str] = []
    L.append(f"# Comps tear-sheet — {s.ticker}")
    L.append("")
    L.append("_Fundamentals: SEC EDGAR (primary). Prices / market cap: yfinance (non-primary)._")
    L.append("")

    # --- Headline: discount to own history ---
    L.append("## Headline — discount to own history")
    L.append("")
    if s.history and s.history.bands:
        L.append("| Multiple | Current | Cheap-vs-own-history %ile | Median | N |")
        L.append("|---|--:|--:|--:|--:|")
        for name, b in s.history.bands.items():
            L.append(f"| {name} | {_f(b.current)} | {_f(b.percentile_cheap, 0)} | {_f(b.median)} | {b.n} |")
        L.append("")
        L.append(_headline_read(result))
    else:
        L.append("_No market cap available — historical bands not computed._")
    L.append("")

    # --- Peer comps ---
    L.append("## Peer comps")
    L.append("")
    L.append("| Ticker | Mkt cap ($M) | P/E | EV/EBITDA | FCF yld | P/TBV | Norm P/E | "
             "P/E %ile own-hist | Implied g | MoS |")
    L.append("|---|--:|--:|--:|--:|--:|--:|--:|--:|--:|")
    for tk, r in result.peer_table.iterrows():
        L.append("| {} | {} | {} | {} | {} | {} | {} | {} | {} | {} |".format(
            tk, _money_m(r.get("market_cap")), _f(r.get("pe")), _f(r.get("ev_ebitda")),
            _pct(r.get("fcf_yield")), _f(r.get("p_tbv")), _f(r.get("normalized_pe")),
            _f(r.get("pe_pctile_own_hist"), 0), _pct(r.get("implied_growth_base")),
            _pct(r.get("margin_of_safety_base")),
        ))
    L.append("")
    if "pe_discount_to_peer_median" in result.peer_table.columns:
        disc = result.peer_table.loc[s.ticker, "pe_discount_to_peer_median"] \
            if s.ticker in result.peer_table.index else float("nan")
        L.append(f"_{s.ticker} P/E vs peer median: **{_pct(disc)}** (negative = cheaper than peers)._")
    L.append("")

    # --- Normalized earnings ---
    L.append("## Normalized earnings (5y mean operating margin × current revenue)")
    L.append("")
    n = s.normalized
    if n is not None:
        L.append(f"- Trailing P/E **{_f(n.trailing_pe)}** · Normalized P/E **{_f(n.normalized_pe)}** "
                 f"· Trough P/E **{_f(n.trough_pe)}** (downside anchor)")
        L.append(f"- Operating margin: {n.n_used}y mean {_pct(n.mean_op_margin)}, "
                 f"trough {_pct(n.trough_op_margin)}")
        L.append(f"- Normalized EPS {_f(n.normalized_eps)} (net income ${_money_m(n.normalized_net_income)}M) "
                 f"· trough EPS {_f(n.trough_eps)}")
    else:
        L.append("_Revenue / operating-margin history insufficient to normalize._")
    L.append("")

    # --- Reverse DCF (downside first) ---
    L.append("## Reverse DCF — FCFF, end-of-year, Gordon terminal")
    L.append("")
    if s.dcf_cases:
        L.append("| Case | WACC | Term g | FCF0 ($M) | Implied g (priced in) | "
                 "Fair value/sh | Price | Margin of safety |")
        L.append("|---|--:|--:|--:|--:|--:|--:|--:|")
        for c in s.dcf_cases:
            ig = f"{c.implied_growth * 100:.1f}%" + ("" if c.implied_converged else "*")
            L.append("| {} | {} | {} | {} | {} | {} | {} | **{}** |".format(
                c.label, _pct(c.wacc), _pct(c.terminal_growth), _money_m(c.fcf0),
                ig, _f(c.fair_value_per_share), _f(c.current_price), _pct(c.margin_of_safety),
            ))
        L.append("")
        L.append("_Bear listed first (downside-first). `*` = solver hit a bracket edge "
                 "(target outside achievable range). Forward fair value uses the conservative "
                 "assumed growth; implied growth is what today's price requires._")
    elif s.is_bank:
        L.append("_N/A — bank; FCF / EBITDA not meaningful. See P/E + P/TBV bands and normalized EPS._")
    else:
        L.append("_Insufficient FCF data for a reverse DCF._")
    L.append("")

    # --- Sources / caveats ---
    L.append("## Sources & caveats")
    L.append("")
    L.append(f"- Fundamentals: **SEC EDGAR XBRL (primary)**. Prices & market cap: "
             f"**yfinance (non-primary)**.")
    L.append(f"- Share basis for historical market caps: **{s.history.share_basis if s.history else 'n/a'}**.")
    L.append(f"- WACC {_pct(result.wacc)}, terminal growth {_pct(result.terminal_growth)}, "
             f"horizon {result.forecast_years}y (config defaults unless overridden).")
    L.append("- FY-end price = unadjusted close on the last trading day ≤ FY-end "
             "(backward as-of, 7-day tolerance). No general split adjustment.")
    L.append("- Own-history percentiles need ≥3 annual observations; thinner history renders n/a.")

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(L) + "\n", encoding="utf-8")

"""Main screener entry point.

For each candidate ticker:
  1. Load cached EDGAR fundamentals (no network).
  2. Optionally fetch current market cap from yfinance (non-primary, tagged).
  3. Compute raw Quality / Valuation / Improving-Fundamentals components.
  4. Apply negative-space filters (Sycamore's revealed style).
  5. Cross-sectionally rank into sub-scores + downside-tilted composite.
  6. Write screener_output.xlsx with sub-scores reported separately.

Network-resilient: if yfinance fails for a ticker, valuation columns are
left blank and the ticker is still included with quality + improving scores.
"""

from __future__ import annotations

import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd

from ..adapters import EdgarProvider, YFinanceProvider
from ..adapters.base import FinancialsFrame
from ..config import cache_dir, load_config
from ..metrics import (
    fcf_margin_fy,
    fcf_yield_fy,
    free_cash_flow_fy,
    gross_margin_fy,
    gross_margin_stability,
    interest_coverage_fy,
    is_bank,
    net_debt_to_ebitda_fy,
    p_tbv_fy,
    pe_fy,
    percentile_vs_history,
    revenue_fy,
    roic_fy,
    rotce_fy,
    ev_ebitda_fy,
)
from ..metrics.profitability import _series, net_income_fy
from ..universe.builder import load_universe
from .scoring import score_universe


# Negative-space thresholds (Sycamore avoids these — exclude or down-rank).
NEG_SPACE_THRESHOLDS = {
    "extreme_pe": 40.0,
    "extreme_ev_ebitda": 25.0,
    "high_leverage_non_bank": 5.0,   # net debt / EBITDA
    "fcf_margin_floor": 0.02,
}

# Volatility-overlay columns appended when run with --vol. These ANNOTATE the
# screen (downside cross-check) and never enter the three-attribute score.
VOL_COLUMNS = [
    "iv_index", "iv_rank", "iv_percentile", "vol_beta", "liquidity_rating",
    "expected_move_30d_pct", "expected_move_earnings_pct", "days_to_earnings",
    "vol_flags", "vol_source",
]


def _join_flags(v) -> str:
    if isinstance(v, list):
        return ", ".join(v)
    return "" if v is None else str(v)


@dataclass
class ScreenerRow:
    ticker: str
    name: str
    gics_sector: str | None
    is_bank: bool
    market_cap: float | None
    # Quality components (raw)
    roic: float | None = None
    rotce: float | None = None
    gross_margin: float | None = None
    gross_margin_stability: float | None = None
    fcf_margin: float | None = None
    interest_coverage: float | None = None
    net_debt_ebitda: float | None = None
    # Valuation components (raw)
    pe: float | None = None
    ev_ebitda: float | None = None
    fcf_yield: float | None = None
    p_tbv: float | None = None
    fcf_yield_pctile_own_history: float | None = None
    # Improving-fundamentals components (raw)
    rev_3yr_cagr: float | None = None
    eps_3yr_cagr: float | None = None
    fcf_3yr_cagr: float | None = None
    rev_growth_acceleration: float | None = None
    # Negative-space flags
    ns_flags: list[str] = field(default_factory=list)
    # Volatility overlay (downside cross-check; NEVER feeds the score/rank).
    iv_index: float | None = None
    iv_rank: float | None = None
    iv_percentile: float | None = None
    vol_beta: float | None = None
    liquidity_rating: float | None = None
    expected_move_30d_pct: float | None = None
    expected_move_earnings_pct: float | None = None
    days_to_earnings: float | None = None
    vol_flags: list[str] = field(default_factory=list)
    vol_source: str = ""
    # Sources contributing to this row
    sources: str = ""
    error: str | None = None


def _last(s: pd.Series) -> float | None:
    s = s.dropna()
    if s.empty:
        return None
    return float(s.iloc[-1])


def _cagr(series: pd.Series, years: int) -> float | None:
    s = series.dropna()
    if len(s) < years + 1:
        return None
    start = s.iloc[-(years + 1)]
    end = s.iloc[-1]
    # A CAGR is only defined for two positive endpoints. A non-positive start
    # OR end (e.g. a swing into a loss year) makes the ratio's fractional power
    # undefined in the reals — guard both so it returns None cleanly rather than
    # producing a NaN + a "invalid value encountered in scalar power" warning.
    if pd.isna(start) or pd.isna(end) or start <= 0 or end <= 0:
        return None
    return float((end / start) ** (1.0 / years) - 1.0)


def _growth_acceleration(series: pd.Series) -> float | None:
    """Last YoY growth minus prior YoY growth. Positive = accelerating."""
    s = series.dropna()
    if len(s) < 3:
        return None
    yoy_last = s.iloc[-1] / s.iloc[-2] - 1.0 if s.iloc[-2] > 0 else None
    yoy_prior = s.iloc[-2] / s.iloc[-3] - 1.0 if s.iloc[-3] > 0 else None
    if yoy_last is None or yoy_prior is None:
        return None
    return float(yoy_last - yoy_prior)


def _negative_space_flags(row: ScreenerRow) -> list[str]:
    flags = []
    if row.fcf_margin is not None and row.fcf_margin < NEG_SPACE_THRESHOLDS["fcf_margin_floor"]:
        flags.append("low_or_neg_fcf")
    if row.pe is not None and row.pe > NEG_SPACE_THRESHOLDS["extreme_pe"]:
        flags.append("extreme_pe")
    if row.ev_ebitda is not None and row.ev_ebitda > NEG_SPACE_THRESHOLDS["extreme_ev_ebitda"]:
        flags.append("extreme_ev_ebitda")
    if (not row.is_bank
        and row.net_debt_ebitda is not None
        and row.net_debt_ebitda > NEG_SPACE_THRESHOLDS["high_leverage_non_bank"]):
        flags.append("high_leverage")
    return flags


def _compute_raw_row(
    ticker: str,
    name: str,
    sector: str | None,
    ff: FinancialsFrame,
    market_cap: float | None,
) -> ScreenerRow:
    bank = is_bank(ff)
    row = ScreenerRow(
        ticker=ticker, name=name, gics_sector=sector,
        is_bank=bank, market_cap=market_cap,
    )
    sources = ["edgar (primary)"]

    # ---- Quality ----
    row.roic = _last(roic_fy(ff))
    row.rotce = _last(rotce_fy(ff)) if bank else None
    row.gross_margin = _last(gross_margin_fy(ff))
    row.gross_margin_stability = gross_margin_stability(ff)
    row.fcf_margin = _last(fcf_margin_fy(ff))
    row.interest_coverage = _last(interest_coverage_fy(ff))
    row.net_debt_ebitda = None if bank else _last(net_debt_to_ebitda_fy(ff))

    # ---- Valuation (requires market_cap) ----
    if market_cap is not None and market_cap > 0:
        row.pe = _last(pe_fy(ff, market_cap))
        row.ev_ebitda = None if bank else _last(ev_ebitda_fy(ff, market_cap))
        row.fcf_yield = _last(fcf_yield_fy(ff, market_cap))
        row.p_tbv = _last(p_tbv_fy(ff, market_cap)) if bank else None
        # "Percentile vs own history" headline: use current FCF yield as a
        # constant against the historical FCF-to-current-mcap series — proxy
        # until Phase 3 lines up historical prices for true bands.
        fcf_series = free_cash_flow_fy(ff)
        if not fcf_series.empty:
            hist_yields = fcf_series / market_cap
            current = row.fcf_yield
            row.fcf_yield_pctile_own_history = percentile_vs_history(
                hist_yields.iloc[:-1], current
            ) if current is not None else None
        sources.append("yfinance (non-primary)")

    # ---- Improving fundamentals ----
    rev = revenue_fy(ff)
    ni = net_income_fy(ff)
    fcf = free_cash_flow_fy(ff)
    row.rev_3yr_cagr = _cagr(rev, 3)
    row.eps_3yr_cagr = _cagr(ni, 3)
    row.fcf_3yr_cagr = _cagr(fcf, 3)
    row.rev_growth_acceleration = _growth_acceleration(rev)

    row.ns_flags = _negative_space_flags(row)
    row.sources = ", ".join(sources)
    return row


def _enrich_rows_with_vol(rows: list[ScreenerRow]) -> str | None:
    """Annotate rows in place with the tastytrade downside vol overlay.

    Pure annotation — never affects scoring or ranking, never raises. Returns a
    human-readable note when the overlay is skipped (no creds / fetch failure),
    else None. Mirrors the screener's existing yfinance resilience.
    """
    from ..adapters import TastytradeProvider
    from ..metrics.volatility import volatility_overlay

    if not TastytradeProvider.available():
        return (
            "vol overlay skipped: set TASTYTRADE_CLIENT_SECRET / "
            "TASTYTRADE_REFRESH_TOKEN to enable (data-only — no positions, no orders)."
        )
    tickers = [r.ticker for r in rows if r.error is None]
    if not tickers:
        return None
    try:
        vf = TastytradeProvider().get_volatility(tickers)
    except Exception as exc:  # noqa: BLE001 — overlay is convenience data
        return f"vol overlay skipped: {exc}"
    for r in rows:
        if r.error is not None:
            continue
        ov = volatility_overlay(vf, r.ticker)
        r.iv_index = ov["iv_index"]
        r.iv_rank = ov["iv_rank"]
        r.iv_percentile = ov["iv_percentile"]
        r.vol_beta = ov["vol_beta"]
        r.liquidity_rating = ov["liquidity_rating"]
        r.expected_move_30d_pct = ov["expected_move_30d_pct"]
        r.expected_move_earnings_pct = ov["expected_move_earnings_pct"]
        r.days_to_earnings = ov["days_to_earnings"]
        r.vol_flags = ov["vol_flags"]
        r.vol_source = ov["vol_source"]
    return None


def _candidate_tickers(
    universe: pd.DataFrame | None,
    explicit: Iterable[str] | None,
    sector: str | None,
    limit: int | None,
) -> list[tuple[str, str, str | None]]:
    """Return [(ticker, name, sector)] tuples for the candidate set."""
    if explicit:
        return [(t.upper(), t.upper(), None) for t in explicit]
    if universe is None:
        cfg = load_config()
        return [(t, t, None) for t in cfg.test_tickers]
    df = universe.copy()
    if sector:
        # case-insensitive contains match on sector string
        mask = df["gics_sector"].fillna("").str.lower().str.contains(sector.lower())
        if sector.lower() == "banks":
            # also accept "Financials" (GICS sector for banks)
            mask = mask | df["gics_sector"].fillna("").str.lower().str.contains("financial")
        df = df[mask]
    if limit:
        df = df.head(limit)
    return [
        (str(r["ticker"]), str(r.get("name", r["ticker"])), r.get("gics_sector"))
        for _, r in df.iterrows()
    ]


# Identifier columns from the scanner that are NOT technical indicators.
_TV_IDENT = {"name", "description", "ticker", "exchange", "type", "subtype",
             "passes_screen", "asof", "source"}


def _tv_divergence(
    composite: pd.Series,
    passes: pd.Series,
    has_tv: pd.Series,
    strong_pctile: float,
) -> pd.Series:
    """Neutral fundamentals-vs-technicals divergence label.

    Derived from the PURE composite score (it never feeds back into it). `has_tv`
    marks rows whose screen membership is known; rows without it get 'no_tv' so a
    failed/partial pull never reads as agreement. NaN composite -> 'n/a'.
    """
    composite = composite.astype(float)
    strong = composite.rank(pct=True) >= strong_pctile
    p = passes.reindex(composite.index).fillna(False).astype(bool)
    htv = has_tv.reindex(composite.index).fillna(False).astype(bool)
    res = pd.Series("n/a", index=composite.index, dtype=object)
    res[~htv] = "no_tv"
    valid = htv & composite.notna()
    res[htv & composite.isna()] = "n/a"
    res[valid & p & strong] = "agree_strong"
    res[valid & ~p & ~strong] = "agree_weak"
    res[valid & ~p & strong] = "diverge_fund_strong_tech_fail"
    res[valid & p & ~strong] = "diverge_fund_weak_tech_pass"
    return res


def _apply_tv_overlay(out: pd.DataFrame, tv_df: pd.DataFrame, tv_cfg) -> pd.DataFrame:
    """Splice the NON-PRIMARY TradingView screen membership + carried indicators
    onto the already-scored frame and add a derived `tv_divergence`.

    Never touches the composite, sub-scores, or rank (CLAUDE.md: bottom-up,
    downside-first). Carried indicators are prefixed `tv_` so provenance is
    obvious and nothing collides with a fundamental column.
    """
    out = out.copy()
    tv = tv_df.copy()
    if "ticker" in tv.columns:
        tv = tv.set_index(tv["ticker"].astype(str).str.upper())
    tv = tv[~tv.index.duplicated(keep="first")]

    passing = set(tv.index)
    passes = pd.Series([t in passing for t in out.index], index=out.index)
    out["passes_screen"] = passes

    ident = _TV_IDENT
    for c in [c for c in tv.columns if c not in ident]:
        out[f"tv_{c}"] = tv[c].reindex(out.index)

    out["tv_asof"] = str(tv["asof"].iloc[0]) if ("asof" in tv.columns and len(tv)) else None

    # Tag `sources` for rows that actually received overlay data (passers). The
    # tag is read from the frame itself, so it is correct for whichever provider
    # (live API or dropped CSV) supplied it.
    src_tag = (
        str(tv["source"].iloc[0])
        if ("source" in tv.columns and len(tv)) else "tradingview (non-primary, technical)"
    )

    def _append_src(cur: object) -> str:
        cur = "" if cur is None or (isinstance(cur, float) and pd.isna(cur)) else str(cur)
        return f"{cur}, {src_tag}" if cur else src_tag
    if "sources" in out.columns:
        out.loc[passes, "sources"] = out.loc[passes, "sources"].apply(_append_src)

    # Once the pull succeeds, membership is known for every row (absence from the
    # screen == fail), so divergence resolves to the four neutral buckets.
    out["tv_divergence"] = _tv_divergence(
        out["composite_score"], passes,
        pd.Series(True, index=out.index),
        float(tv_cfg.divergence_strong_pctile),
    )
    return out


def _tv_provider(mode: str):
    """Pick the technical-overlay provider for the configured mode.

    'trend' computes a transparent Bull/Neutral/Bear regime in Python from the
    prices we already pull (no TradingView dependency); 'csv' reads a dropped
    CSV (e.g. a custom Pine indicator's flagged tickers); 'api' replicates a
    built-in TradingView Stock Screener live. All implement
    TechnicalScreenProvider, so the overlay code downstream is identical.
    """
    m = str(mode).lower()
    if m == "trend":
        from ..adapters.trend_regime import TrendRegimeProvider
        return TrendRegimeProvider()
    if m == "csv":
        from ..adapters.csv_screen import CsvScreenProvider
        return CsvScreenProvider()
    from ..adapters.tradingview import TradingViewProvider
    return TradingViewProvider()


def run_screener(
    tickers: Iterable[str] | None = None,
    sector: str | None = None,
    limit: int | None = None,
    output_path: Path | str | None = None,
    skip_market_cap: bool = False,
    hard_exclude_neg_space: bool = False,
    with_vol: bool = False,
    tv_overlay: bool = False,
    refresh_tv: bool = False,
    with_prediction_overlay: bool = False,
) -> pd.DataFrame:
    """Run the three-attribute screen and write screener_output.xlsx.

    Negative-space handling (CLAUDE.md principle 2 — decompose, never hide):
    by default, names that trip a negative-space flag are KEPT in the output
    with all three sub-scores visible, just sorted to the bottom with their
    flags shown. Set hard_exclude_neg_space=True to drop them from the
    ranking entirely (composite_rank = NaN).
    """
    edgar = EdgarProvider()
    yfin = None if skip_market_cap else YFinanceProvider()

    universe = load_universe()
    candidates = _candidate_tickers(universe, tickers, sector, limit)
    if not candidates:
        raise ValueError(
            "No candidate tickers. Either pass tickers explicitly, "
            "run `build-universe` first, or set test_tickers in config.yaml."
        )

    rows: list[ScreenerRow] = []
    for ticker, name, sec in candidates:
        try:
            ff = edgar.get_financials(ticker)
        except Exception as exc:  # noqa: BLE001
            rows.append(ScreenerRow(
                ticker=ticker, name=name, gics_sector=sec, is_bank=False,
                market_cap=None, error=f"financials load failed: {exc}",
            ))
            continue
        mcap = None
        if yfin is not None:
            try:
                mcap = yfin.get_market_cap(ticker)
            except Exception:  # noqa: BLE001 — yfinance is intentionally flaky
                mcap = None
            # Fallback: price × EDGAR shares. Keeps share count primary-source
            # per CLAUDE.md and survives yfinance market-cap outages.
            if mcap is None:
                try:
                    price = yfin.get_current_price(ticker)
                    shares = _last(_series(ff, "SharesOutstanding"))
                    if price and shares:
                        mcap = float(price) * float(shares)
                except Exception:  # noqa: BLE001
                    mcap = None
        rows.append(_compute_raw_row(ticker, name, sec, ff, mcap))

    # Volatility overlay (downside cross-check) — annotation only, never scored.
    vol_note = _enrich_rows_with_vol(rows) if with_vol else None

    raw = pd.DataFrame([r.__dict__ for r in rows]).set_index("ticker")
    # Stringify list columns for xlsx output.
    raw["ns_flags"] = raw["ns_flags"].apply(_join_flags)
    raw["negative_space"] = raw["ns_flags"].fillna("").str.len() > 0
    if with_vol and vol_note is None:
        raw["vol_flags"] = raw["vol_flags"].apply(_join_flags)
    else:
        # No vol data produced — drop the empty overlay columns entirely.
        raw = raw.drop(columns=[c for c in VOL_COLUMNS if c in raw.columns])

    # Score ALL names — sub-scores stay honest even for flagged names so the
    # three-axis decomposition is never hidden (CLAUDE.md principle 2). The
    # negative-space policy only affects the RANK, not the sub-scores.
    result = score_universe(raw)
    score_cols = [c for c in result.df.columns
                  if c.endswith("_score") or c.endswith("_pctile")]
    out = raw.join(result.df[score_cols], how="left")

    if hard_exclude_neg_space:
        # Flagged names removed from the ranking entirely (rank = NaN), but
        # kept in the output with sub-scores + flags for transparency.
        clean = out.loc[~out["negative_space"], "composite_score"]
        out["composite_rank"] = clean.rank(ascending=False, method="min")
        out = out.sort_values(
            ["negative_space", "composite_score"],
            ascending=[True, False], na_position="last",
        )
    else:
        # Keep flagged names but down-rank: clean names first (by composite),
        # flagged names after (by their own composite). composite_score stays
        # the honest weighted blend — only the rank encodes the down-ranking.
        out = out.sort_values(
            ["negative_space", "composite_score"],
            ascending=[True, False], na_position="last",
        )
        out["composite_rank"] = range(1, len(out) + 1)
        out.loc[out["composite_score"].isna(), "composite_rank"] = np.nan

    # ---- TradingView technical overlay (NON-PRIMARY context) ----
    # Quarantined per CLAUDE.md (bottom-up only): spliced in AFTER the composite
    # and rank are final, so it is structurally unable to reach score_universe.
    # Any failure (disabled in config, network/CSV/creds absent) degrades to the
    # full fundamental output with the TA columns absent + a clear, actionable
    # skip note — mirrors the vol overlay and the pipeline's missing-universe
    # message; never a silent no-op and never a bare traceback.
    cfg = load_config()
    tv_note: str | None = None
    if tv_overlay:
        if not cfg.tradingview.enabled:
            tv_note = (
                "tv overlay skipped: set tradingview.enabled: true in config.yaml "
                "(and choose mode: trend | csv | api) to enable the TradingView "
                "technical overlay. The fundamental screen is unaffected."
            )
        else:
            try:
                tv_df = _tv_provider(cfg.tradingview.mode).get_screen(
                    tickers=list(out.index), refresh=refresh_tv
                )
                out = _apply_tv_overlay(out, tv_df, cfg.tradingview)
            except Exception as exc:  # noqa: BLE001 — overlay must never break the screen
                tv_note = f"tv overlay skipped: {exc}"
                print(f"[tv-overlay] {tv_note}", file=sys.stderr)

    # Optional NON-PRIMARY prediction-market annotations. Joined as separate
    # event_* columns — never merged into composite_score/rank (CLAUDE.md:
    # the prediction lens is walled off from every scoring path). Returns a clear
    # skip note when nothing surfaced (egress blocked / no confirmed mappings).
    prediction_note: str | None = None
    if with_prediction_overlay:
        out, prediction_note = _attach_prediction_overlay(out)

    # Reorder columns: identity, sub-scores prominent, then components and
    # sources, and finally the NON-PRIMARY TradingView overlay columns — kept to
    # the RIGHT of the downside flags so margin-of-safety stays most prominent.
    front = [
        "name", "gics_sector", "is_bank", "market_cap",
        "composite_rank", "composite_score",
        "q1_quality_score", "q2_valuation_score", "q3_improving_score",
        "negative_space", "ns_flags",
        # Volatility overlay (downside cross-check) — kept prominent, near the
        # scores, per CLAUDE.md (downside at least as visible as upside).
        "iv_rank", "iv_percentile", "iv_index",
        "expected_move_30d_pct", "expected_move_earnings_pct",
        "days_to_earnings", "vol_beta", "liquidity_rating", "vol_flags",
    ]
    front = [c for c in front if c in out.columns]
    tv_order = ["passes_screen", "tv_divergence"]
    tv_indicators = sorted(
        c for c in out.columns if c.startswith("tv_") and c not in ("tv_divergence", "tv_asof")
    )
    tv_tail = [c for c in (tv_order + tv_indicators + ["tv_asof"]) if c in out.columns]
    rest = [c for c in out.columns if c not in front and c not in tv_tail]
    out = out[front + rest + tv_tail]

    # Surface the overlay notes (e.g. "skipped: no creds") off the frame so the
    # CLI can print them without coupling library code to typer. Each is None
    # when its overlay ran clean or was not requested.
    out.attrs["vol_note"] = vol_note
    out.attrs["tv_note"] = tv_note
    out.attrs["prediction_note"] = prediction_note

    output_path = Path(output_path) if output_path else (cache_dir() / "screener_output.xlsx")
    _write_screener_xlsx(out, output_path)
    return out


def _write_screener_xlsx(df: pd.DataFrame, path: Path) -> None:
    """Single-sheet xlsx. Header row frozen, basic auto-width."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "screener"

    # Header
    cols = ["ticker"] + list(df.columns)
    for c, v in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=v)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
        cell.alignment = Alignment(horizontal="left")

    # Rows
    for r, (ticker, row) in enumerate(df.iterrows(), start=2):
        ws.cell(row=r, column=1, value=ticker)
        for c, col in enumerate(df.columns, start=2):
            v = row[col]
            if isinstance(v, float) and (pd.isna(v) or np.isinf(v)):
                v = None
            ws.cell(row=r, column=c, value=v)

    ws.freeze_panes = "B2"
    # Rough column-width pass.
    for i, col in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, min(28, len(str(col)) + 4))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _attach_prediction_overlay(out: pd.DataFrame) -> tuple[pd.DataFrame, str | None]:
    """Join NON-PRIMARY prediction-market annotations (event_* columns) onto the
    screener output. Returns ``(annotated_frame, note)``.

    Best-effort and isolated: it never raises into the screen (Polymarket egress
    is often blocked), and it NEVER feeds composite_score/rank — a separate lens
    per CLAUDE.md. `event_contradiction` flags an otherwise high-ranked, clean
    name that nonetheless carries a material market-implied RISK — exactly the
    signal the bottom-up screen can't see. `note` is a clear, actionable message
    when the overlay surfaced nothing (egress blocked / no confirmed mappings),
    else None — so a skip is never a silent no-op or a bare traceback.
    """
    try:
        from ..adapters import PolymarketProvider
        from ..prediction_markets.overlay import build_overlay, screener_annotations

        overlay = build_overlay(
            PolymarketProvider(), tickers=list(out.index), confirmed_only=True
        )
        ann = screener_annotations(overlay)
        out = out.join(ann, how="left")

        cp = load_config().prediction.contradiction_prob
        n = len(out)
        prob = pd.to_numeric(out.get("event_top_prob"), errors="coerce")
        rank = pd.to_numeric(out.get("composite_rank"), errors="coerce")
        neg = out.get("negative_space")
        neg = pd.Series(False, index=out.index) if neg is None else neg.fillna(False).astype(bool)
        out["event_contradiction"] = (
            (prob.fillna(0.0) >= cp)
            & (out.get("event_read_through") == "RISK")
            & (~neg)
            & (rank.fillna(n) <= (n / 2.0))
        )
        note = None
        if prob.notna().sum() == 0:
            note = (
                "prediction overlay: no confirmed market mappings surfaced. Run "
                "`prediction-discover` and set confirmed=True in "
                "data/raw/prediction_markets.csv (needs Polymarket egress). The "
                "event_* columns are blank; the fundamental screen is unaffected."
            )
        return out, note
    except Exception as exc:  # noqa: BLE001 — overlay must never break the screen
        out["event_overlay_error"] = str(exc)
        return out, f"prediction overlay skipped: {exc}"

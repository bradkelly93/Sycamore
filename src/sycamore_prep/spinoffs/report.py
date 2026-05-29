"""Writers for the spin-off tracker: a multi-tab xlsx + a markdown tear-sheet.

Reuses the screener/comps openpyxl conventions (bold frozen header, auto-width,
NaN/None -> blank). Downside-first: a dedicated downside-flags tab/section sits
ahead of the upside; the three attributes stay separate; missing data renders
"pending", never a silent zero; every row carries a source tag + a filing link.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd

# Columns surfaced on the headline "tracker" tab, in downside-first order.
_TRACKER_COLS = [
    "spinco_name", "spinco_ticker", "parent_ticker", "status",
    "downside_flags", "review_flags",
    "first_form10_date", "latest_amendment_date", "amendment_count",
    "distribution_ratio", "distribution_ratio_source", "record_date", "distribution_date",
    "q1_better_business", "q2_valuation_disparity", "q3_improving_fundamentals",
    "sic", "primary_filing_url", "sources", "error",
]
_DOWNSIDE_COLS = [
    "spinco_name", "spinco_ticker", "status", "has_financials",
    "net_debt_ebitda", "interest_coverage", "negative_equity",
    "fcf_margin", "revenue_cagr_3y", "downside_flags", "review_flags",
    "primary_filing_url",
]
_FIN_COLS = [
    "spinco_name", "spinco_ticker", "has_financials",
    "net_debt_ebitda", "interest_coverage", "fcf", "fcf_margin", "roic",
    "revenue_cagr_3y", "ebitda_is_proxy",
    "q1_better_business", "q3_improving_fundamentals", "sources",
]


def _safe(v):
    """Coerce to something openpyxl can store; NaN/Inf/None -> None (blank)."""
    if isinstance(v, np.generic):
        v = v.item()
    if v is None:
        return None
    if isinstance(v, float) and (pd.isna(v) or np.isinf(v)):
        return None
    if isinstance(v, (int, float, str, bool)):
        return v
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    return str(v)


def _join(v) -> str:
    return ", ".join(v) if isinstance(v, (list, tuple)) else ("" if v is None else str(v))


def records_to_df(records: list) -> pd.DataFrame:
    """One tidy row per spin-off. List columns are joined to strings; numeric
    gaps stay NaN (blank); the three q-attributes stay 'pending' strings."""
    rows = []
    for r in records:
        urls = r.filing_urls or []
        rows.append({
            "spinco": r.spinco_ticker or r.spinco_name or r.spinco_cik or "?",
            "spinco_name": r.spinco_name,
            "spinco_ticker": r.spinco_ticker,
            "spinco_cik": r.spinco_cik,
            "parent_name": r.parent_name,
            "parent_ticker": r.parent_ticker,
            "parent_cik": r.parent_cik,
            "status": r.status,
            "form_type": r.form_type,
            "first_form10_date": r.first_form10_date,
            "latest_amendment_date": r.latest_amendment_date,
            "amendment_count": r.amendment_count,
            "distribution_ratio": r.distribution_ratio,
            "distribution_ratio_source": r.distribution_ratio_source,
            "record_date": r.record_date,
            "distribution_date": r.distribution_date,
            "has_financials": r.has_financials,
            "net_debt_ebitda": r.net_debt_ebitda,
            "interest_coverage": r.interest_coverage,
            "fcf": r.fcf,
            "fcf_margin": r.fcf_margin,
            "roic": r.roic,
            "revenue_cagr_3y": r.revenue_cagr_3y,
            "negative_equity": r.negative_equity,
            "ebitda_is_proxy": r.ebitda_is_proxy,
            "q1_better_business": r.q1_better_business,
            "q2_valuation_disparity": r.q2_valuation_disparity,
            "q3_improving_fundamentals": r.q3_improving_fundamentals,
            "downside_flags": _join(r.downside_flags),
            "review_flags": _join(r.review_flags),
            "sic": r.sic,
            "n_filings": len(r.accessions or []),
            "primary_filing_url": r.information_statement_url or (urls[-1] if urls else None),
            "sources": r.sources,
            "error": r.error,
        })
    df = pd.DataFrame(rows)
    if not df.empty:
        df = df.set_index("spinco")
    return df


def _filings_df(records: list) -> pd.DataFrame:
    rows = []
    for r in records:
        key = r.spinco_ticker or r.spinco_name or r.spinco_cik or "?"
        accs = r.accessions or []
        urls = r.filing_urls or []
        for i in range(max(len(accs), len(urls))):
            rows.append({
                "spinco": key,
                "accession": accs[i] if i < len(accs) else None,
                "filing_url": urls[i] if i < len(urls) else None,
            })
    df = pd.DataFrame(rows, columns=["spinco", "accession", "filing_url"])
    return df.set_index("spinco") if not df.empty else df


def _sources_df(mode: str) -> pd.DataFrame:
    items = {
        "fundamentals": "SEC EDGAR XBRL companyfacts (primary)",
        "filings_and_discovery": "SEC EDGAR submissions + full-text search (primary)",
        "distribution_ratio_dates": "parsed from the 10-12B information statement when "
                                    "unambiguous (derived); else 'pending — read Form 10'",
        "note_pending": "A freshly-registered SpinCo has no XBRL until its first 10-K; "
                        "leverage/quality render 'pending', never zero.",
        "note_three_attribute": "Q1 better-business / Q2 valuation-disparity / Q3 improving "
                                "are reported SEPARATELY; there is no single composite.",
        "note_downside_first": "Downside flags + 'review' prompts are surfaced at least as "
                               "prominently as upside (CLAUDE.md #1).",
        "note_review": "'review' flags require reading the linked Form 10 (pension / "
                       "litigation / debt transfer, reason for the spin).",
        "mode": mode,
    }
    df = pd.DataFrame({"value": items})
    df.index.name = "key"
    return df


def _df_to_sheet(ws, df: pd.DataFrame, index_label: str) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    if df is None or df.empty:
        ws.cell(row=1, column=1, value="(no data)")
        return
    cols = [index_label] + list(df.columns)
    for c, v in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=str(v))
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
        cell.alignment = Alignment(horizontal="left")
    for r, (idx, rowv) in enumerate(df.iterrows(), start=2):
        ws.cell(row=r, column=1, value=_safe(idx))
        for c, col in enumerate(df.columns, start=2):
            ws.cell(row=r, column=c, value=_safe(rowv[col]))
    ws.freeze_panes = "B2"
    for i, col in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(
            12, min(40, len(str(col)) + 4)
        )


def write_tracker_xlsx(result, path: Path) -> None:
    from openpyxl import Workbook

    df = result.df
    tabs = [
        ("tracker", df.reindex(columns=[c for c in _TRACKER_COLS if c in df.columns]) if not df.empty else df, "spinco"),
        ("downside_flags", df.reindex(columns=[c for c in _DOWNSIDE_COLS if c in df.columns]) if not df.empty else df, "spinco"),
        ("spinco_financials", df.reindex(columns=[c for c in _FIN_COLS if c in df.columns]) if not df.empty else df, "spinco"),
        ("filings", _filings_df(result.records), "spinco"),
        ("sources", _sources_df(result.mode), "key"),
    ]
    wb = Workbook()
    for i, (title, tdf, idx_label) in enumerate(tabs):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        _df_to_sheet(ws, tdf, idx_label)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _fmt(v, pending_ok: bool = True) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "pending" if pending_ok else "n/a"
    return str(v)


def _pct(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "pending"
    return f"{float(v) * 100:.1f}%"


def _record_section(r) -> list[str]:
    L: list[str] = []
    head = r.spinco_name or r.spinco_ticker or r.spinco_cik or "Unknown SpinCo"
    L.append(f"## {head}" + (f" ({r.spinco_ticker})" if r.spinco_ticker else ""))
    L.append("")
    L.append(f"- **Parent:** {_fmt(r.parent_name)} ({_fmt(r.parent_ticker)}) · "
             f"**Status:** {r.status}")
    L.append(f"- **Form 10:** first {_fmt(r.first_form10_date)}, "
             f"latest amendment {_fmt(r.latest_amendment_date)} "
             f"({r.amendment_count} amendment(s))")
    L.append(f"- **Distribution:** ratio {_fmt(r.distribution_ratio)}"
             + (f" [{r.distribution_ratio_source}]" if r.distribution_ratio_source else "")
             + f" · record {_fmt(r.record_date)} · distributed {_fmt(r.distribution_date)}")
    L.append("")
    # Downside first.
    L.append("**Downside flags:** "
             + (", ".join(r.downside_flags) if r.downside_flags else "_none computed_"))
    if r.review_flags:
        link = r.information_statement_url or (r.filing_urls[-1] if r.filing_urls else None)
        L.append("**Review (read Form 10):** " + "; ".join(r.review_flags)
                 + (f" — [Form 10]({link})" if link else ""))
    L.append("")
    # Three attributes, separate.
    L.append("**Three-attribute read (kept separate):**")
    L.append(f"- Q1 better business: {_fmt(r.q1_better_business)}")
    L.append(f"- Q2 valuation disparity: {_fmt(r.q2_valuation_disparity)}")
    L.append(f"- Q3 improving fundamentals: {_fmt(r.q3_improving_fundamentals)}")
    if r.has_financials:
        L.append("")
        L.append(f"- net debt/EBITDA {_fmt(r.net_debt_ebitda)} · interest coverage "
                 f"{_fmt(r.interest_coverage)} · FCF margin {_pct(r.fcf_margin)} · "
                 f"rev 3y CAGR {_pct(r.revenue_cagr_3y)}"
                 + (" · _EBITDA is EBIT proxy_" if r.ebitda_is_proxy else ""))
    L.append("")
    L.append(f"_Sources: {r.sources}._"
             + (f" _Note: {r.error}_" if r.error else ""))
    L.append("")
    return L


def write_tracker_md(result, path: Path) -> None:
    L: list[str] = []
    L.append(f"# Spin-off tracker — {result.mode}")
    L.append("")
    L.append("_Filings & fundamentals: SEC EDGAR (primary). Distribution ratio/dates: "
             "parsed from the 10-12B when unambiguous (derived), else pending. "
             "Three attributes reported separately; downside surfaced first._")
    L.append("")
    if not result.records:
        L.append("_No spin-offs found for this query._")
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("\n".join(L) + "\n", encoding="utf-8")
        return

    if result.mode == "scan":
        L.append("## Candidates (recent Form 10 / 10-12B registrations)")
        L.append("")
        L.append("| SpinCo | Ticker | Status | First 10-12B | Amendments | SIC | Form 10 |")
        L.append("|---|---|---|---|--:|---|---|")
        for r in result.records:
            link = r.filing_urls[-1] if r.filing_urls else None
            L.append("| {} | {} | {} | {} | {} | {} | {} |".format(
                _fmt(r.spinco_name), _fmt(r.spinco_ticker), r.status,
                _fmt(r.first_form10_date), r.amendment_count, _fmt(r.sic),
                f"[link]({link})" if link else "n/a",
            ))
        L.append("")
        L.append("_Scan is broad discovery (metadata + status). Run "
                 "`spinoffs track <PARENT> --spinco <TICKER>` for leverage/quality flags._")
    else:
        for r in result.records:
            L.extend(_record_section(r))

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(L) + "\n", encoding="utf-8")
